#!/usr/bin/env python3
"""Importa drilldowns GSC (xlsx) → scripts/gsc-coverage-urls.csv para comprobar-indexacion-gestoria."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print('Instala openpyxl: pip install openpyxl', file=sys.stderr)
    raise

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'scripts' / 'gsc-coverage-urls.csv'

GESTORIA_HINTS = (
    '/gestoria',
    '/contratos-inmobiliarios',
    '/contrato-alquiler',
    '/contrato-arras',
)


def url_relevant(url: str) -> bool:
    return any(h in url for h in GESTORIA_HINTS)


def import_files(paths: list[Path]) -> int:
    rows = ['issue,url,category,prefix']
    for path in paths:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        meta = {r[0]: r[1] for r in wb['Metadatos'].iter_rows(values_only=True) if r[0]}
        issue = str(meta.get('Incidencia', 'Unknown'))
        for r in wb['Tabla'].iter_rows(values_only=True):
            u = r[0]
            if not u or u == 'URL':
                continue
            url = str(u).split()[0]
            if url_relevant(url):
                rows.append(f'{issue},{url},gestoria,')
        wb.close()
    OUT.write_text('\n'.join(rows) + '\n', encoding='utf-8')
    return len(rows) - 1


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('xlsx', nargs='+', type=Path, help='Archivos Coverage-Drilldown*.xlsx')
    args = p.parse_args()
    n = import_files(args.xlsx)
    print(f'Escritas {n} filas en {OUT}')


if __name__ == '__main__':
    main()
